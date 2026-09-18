"""Build the runtime JCR/CAS quality index from the supplied JCR workbook.

The generated JSON is intentionally git-ignored: it is derived from a licensed
source workbook and is mounted as runtime data for the public service.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


def text(value: object) -> str:
    return str(value or "").strip()


def normalized(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", text(value).lower())


def normalized_issn(value: object) -> str:
    return re.sub(r"[^0-9x]", "", text(value).lower())


def column_index(headers: list[object], name: str) -> int:
    for index, header in enumerate(headers):
        if text(header).casefold() == name.casefold():
            return index
    raise ValueError(f"Required JCR column is missing: {name}")


def cas_warning_sets(path: Path) -> tuple[set[str], set[str]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if payload.get("list_year") != 2025:
        raise ValueError("Only the 2025 CAS warning list is accepted.")
    title_set: set[str] = set()
    issn_set: set[str] = set()
    for item in payload.get("journals", []):
        title_set.add(normalized(item.get("title")))
        for value in item.get("issn", []):
            if number := normalized_issn(value):
                issn_set.add(number)
    return title_set - {""}, issn_set


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--cas-list", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    warnings_by_title, warnings_by_issn = cas_warning_sets(args.cas_list)
    workbook = load_workbook(args.xlsx, read_only=True, data_only=True)
    worksheet = workbook.active
    headers = list(next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True)))
    title_col = 0
    issn_col = column_index(headers, "ISSN")
    eissn_col = column_index(headers, "EISSN")
    quartile_col = column_index(headers, "分区")
    five_year_col = column_index(headers, "5 Year JIF")

    records: list[dict] = []
    unique: set[tuple[str, str, str, str, float]] = set()
    skipped = 0
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        journal = text(row[title_col])
        quartile = text(row[quartile_col]).upper()
        try:
            five_year_jif = float(row[five_year_col])
        except (TypeError, ValueError):
            skipped += 1
            continue
        if not journal or quartile not in {"Q1", "Q2", "Q3", "Q4"}:
            skipped += 1
            continue
        issn = text(row[issn_col]) or None
        eissn = text(row[eissn_col]) or None
        identity = (normalized(journal), normalized_issn(issn), normalized_issn(eissn), quartile, five_year_jif)
        if identity in unique:
            continue
        unique.add(identity)
        cas_warning = normalized(journal) in warnings_by_title or bool(
            {normalized_issn(issn), normalized_issn(eissn)} & warnings_by_issn
        )
        records.append(
            {
                "journal": journal,
                "issn": issn,
                "eissn": eissn,
                "jcr_release_year": 2026,
                "jcr_quartile": quartile,
                "five_year_jif": five_year_jif,
                "cas_warning_2025": cas_warning,
                "source": "2026年JCR期刊分区信息.xlsx + CAS 2025 warning list",
            }
        )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"wrote={len(records)} skipped={skipped} cas_flagged={sum(item['cas_warning_2025'] for item in records)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
